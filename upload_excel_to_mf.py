import configparser
from datetime import datetime
import logging
import os
import sys
import time

import openpyxl
import pandas as pd
from selenium import webdriver
from selenium.webdriver.common.by import By
from selenium.webdriver.support import expected_conditions as EC
from selenium.webdriver.support.ui import WebDriverWait

# 設定ファイル
CONFIG_FILE = ".\config.ini"

# Edgeサービス
EDGE_SERVICE = webdriver.EdgeService(executable_path=".\msedgedriver.exe")

# ログ設定
LOGGER = logging.getLogger(__name__)
LOGGER.setLevel(logging.INFO)
LOGGER_FORMATTER = logging.Formatter("%(asctime)s [%(name)s] %(levelname)s %(message)s")

LOG_FOLDER=".\log"
if not os.path.exists(LOG_FOLDER):
    os.mkdir(LOG_FOLDER)
today_str = datetime.today().strftime("%Y%m%d")
LOG_HANDLER = logging.FileHandler(os.path.join(LOG_FOLDER, f"log_{today_str}.log"), encoding="utf-8")
LOG_HANDLER.setFormatter(LOGGER_FORMATTER)
LOGGER.addHandler(LOG_HANDLER)

LOG_CONSOLE_HANDLER = logging.StreamHandler(sys.stdout)
LOG_CONSOLE_HANDLER.setFormatter(LOGGER_FORMATTER)
LOGGER.addHandler(LOG_CONSOLE_HANDLER)

def read_config(filename: str) -> configparser.ConfigParser:
    """
    設定ファイルを読み込む関数

    Args:
        filename (str): 読み込む設定ファイルの名前

    Returns:
        ConfigParser: 設定データを含む ConfigParser オブジェクト

    """
    config = configparser.ConfigParser()
    config.read(filename)
    return config

def get_setting(config: configparser.ConfigParser, section: str, option: str) -> str:
    """
    ConfigParser オブジェクトから設定値を取得する関数

    Args:
        config (ConfigParser): 設定データを含む ConfigParser オブジェクト
        section (str): 設定のセクション名
        option (str): 指定したセクション内のオプション名

    Returns:
        str: 指定したセクション内のオプションの値
    """
    return config.get(section, option)

def read_excel_table(file_path: str, table_name: str) -> pd.DataFrame:
    """
    指定されたExcelシートの特定のテーブルからデータを読み取る関数
    
    Args:
        file_path (str): Excelファイルのパス
        table_name (str): 読み取るテーブルの名前
    
    Returns:
        pd.DataFrame: 読み取ったデータを含むDataFrame
    """
    
    def get_table_range(table_ref: str) -> tuple:
        """
        テーブルの参照範囲を取得する関数
        
        Args:
            table_ref (str): テーブルの参照範囲（例: 'A1:D10'）
        
        Returns:
            tuple: テーブルの最小行、最大行、最小列、最大列を含むタプル
        """
        min_row, max_row = table_ref.split(":")[0][1:], table_ref.split(":")[1][1:]
        min_col, max_col = table_ref.split(":")[0][0], table_ref.split(":")[1][0]
        
        min_col_idx = openpyxl.utils.column_index_from_string(min_col)
        max_col_idx = openpyxl.utils.column_index_from_string(max_col)
        
        return int(min_row), int(max_row), min_col_idx, max_col_idx
    
    workbook = openpyxl.load_workbook(file_path, data_only=True)
    worksheet = workbook.active
    table = worksheet.tables[table_name]
    
    min_row, max_row, min_col_idx, max_col_idx = get_table_range(table.ref)
    data = list(worksheet.iter_rows(min_row=min_row, max_row=max_row, min_col=min_col_idx, max_col=max_col_idx, values_only=True))
    
    dataframe = pd.DataFrame(data)
    dataframe.columns = dataframe.iloc[0]
    dataframe = dataframe[1:]
    
    return dataframe

def setup_webdriver(service: webdriver.EdgeService) -> webdriver.Edge:
    """Edge WebDriverのセットアップを行う

    Args:
        service (webdriver.EdgeService): Edge WebDriverのサービス

    Returns:
        webdriver.Edge: セットアップされたEdge WebDriver
    """
    driver = webdriver.Edge(service=service)
    driver.implicitly_wait(10)
    return driver

def perform_login(driver: webdriver.Edge, signin_url: str, user: str, password: str) -> None:
    """ログインを行う

    Args:
        driver (webdriver.Edge): Edge WebDriver
        signin_url (str): サインインURL
        user (str): ユーザー名
        password (str): パスワード
    """
    driver.get(signin_url)
    driver.implicitly_wait(10)

    elem = driver.find_element(By.NAME, "mfid_user[email]")
    elem.clear()
    elem.send_keys(user)
    elem.submit()
    driver.implicitly_wait(10)

    elem = driver.find_element(By.NAME, "mfid_user[password]")
    elem.clear()
    elem.send_keys(password)
    elem.submit()

def wait_for_element(driver: webdriver.Edge, by: str, identifier: str, timeout: int = 10) -> None:
    """指定した要素が存在するまで待機する

    Args:
        driver (webdriver.Edge): Edge WebDriver
        by (str): 要素の検索方法
        identifier (str): 要素の識別子
        timeout (int, optional): タイムアウト時間. Defaults to 10.
    """
    WebDriverWait(driver, timeout).until(EC.presence_of_element_located((by, identifier)))

def input_data(driver: webdriver.Edge, row: pd.Series, wallet_xpath: str) -> None:
    """入出金登録する

    Args:
        driver (webdriver.Edge): Edge WebDriver
        row (pd.Series): 登録するデータの行
        wallet_xpath (str): 入出金元のXPATH
    """
    # 行にnullが含まれる場合、登録をスキップする
    if row.isnull().any():
        LOGGER.error(f"行に空値が含まれているため、登録をスキップします。行データ：\r\n{row}")
        return
    
    # 金額が正の場合、収入とみなす
    if int(row["Amount"]) > 0:
        driver.find_element(By.CLASS_NAME, "plus-payment").click()
    
    # 金額の入力
    elem = driver.find_element(By.ID, "appendedPrependedInput")
    elem.clear()
    price = abs(int(row["Amount"]))
    elem.send_keys(price)

    # ウォレットの入力
    driver.find_element(By.ID, "user_asset_act_sub_account_id_hash").click()
    time.sleep(1)
    driver.find_element(By.XPATH, wallet_xpath).click()
    time.sleep(1)

    # 大カテゴリの入力
    driver.find_element(By.ID, "js-large-category-selected").click()
    time.sleep(1)
    large_category = row["Large Category"]
    driver.find_element(By.XPATH, f"//a[text()='{large_category}' and @class='l_c_name']").click()
    time.sleep(1)

    # 中カテゴリの入力
    driver.find_element(By.ID, "js-middle-category-selected").click()
    time.sleep(1)
    middle_category = row["Middle Category"]
    driver.find_element(By.XPATH, f"//a[text()='{middle_category}' and @class='m_c_name']").click()
    time.sleep(1)

    # 内容の入力
    elem = driver.find_element(By.ID, "js-content-field")
    elem.clear()
    content = row["Content"]
    elem.send_keys(content)
    
    # 日付の入力
    elem = driver.find_element(By.ID, "updated-at")
    elem.clear()
    date = row['Date'].strftime("%Y/%m/%d")
    elem.send_keys(date)

    # 保存
    driver.find_element(By.ID, "submit-button").click()

    # 確認ボタンの待機とクリック
    wait_for_element(driver, By.ID, "confirmation-button")
    driver.find_element(By.ID, "confirmation-button").click()
    LOGGER.info(f"入出金登録を行いました。日付：{date} 大カテゴリ：{large_category} 中カテゴリ：{middle_category} 備考：{content} 金額：{price}")

    # 次のサイクルのフォーム準備を待機
    wait_for_element(driver, By.ID, "submit-button")
    wait_for_element(driver, By.CLASS_NAME, "plus-payment")

# プログラムのエントリーポイント
if __name__ == "__main__":
    try:
        LOGGER.info("処理開始")
        
        # 設定ファイルを読み込む
        LOGGER.info("設定ファイル読込処理を開始しました。")
        config = read_config(CONFIG_FILE)
        input_file = get_setting(config, "settings", "input_file")
        table_name = get_setting(config, "settings", "table_name")
        user = get_setting(config, "settings", "user")
        password = get_setting(config, "settings", "password")
        signin_url = get_setting(config, "settings", "signin_url")
        input_url = get_setting(config, "settings", "input_url")
        wallet_xpath = get_setting(config, "settings", "wallet_xpath")
        LOGGER.info("設定ファイル読込処理を完了しました。")
        
        try:
            # WebDriverのセットアップ
            LOGGER.info("WebDriverのセットアップを開始しました。")
            driver = setup_webdriver(EDGE_SERVICE)
            LOGGER.info("WebDriverのセットアップが完了しました。")
            
            # ログインの実行
            LOGGER.info(f"Money Forward MEへのログインを開始しました。")
            perform_login(driver, signin_url, user, password)

            # データ入力画面への遷移とボタンの待機
            driver.get(input_url)
            wait_for_element(driver, By.ID, "submit-button")
            LOGGER.info("Money Forward MEへのログインが完了しました。")

            # Excelファイルの読み込み
            LOGGER.info(f"Excelファイルの読み込みを開始しました。パス：{input_file}")
            df = read_excel_table(input_file, table_name)
            LOGGER.info(f"Excelファイルの読み込みが完了しました。")
            
            # 入出金登録
            LOGGER.info("Money Forward MEのPayPayへの入出金登録を開始しました。")
            for index, row in df.iterrows():
                input_data(driver, row, wallet_xpath)
            LOGGER.info("Money Forward MEのPayPayへの入出金登録が完了しました。")
        finally:
            driver.quit()
            LOGGER.info("WebDriverを終了しました。")
        
        LOGGER.info("処理終了")
    except:
        LOGGER.exception("異常終了しました。")
